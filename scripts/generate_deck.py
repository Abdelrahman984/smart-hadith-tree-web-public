import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_pptx):
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_COLOR = RGBColor(11, 19, 43)        # #0B132B Dark Navy
    CARD_BG = RGBColor(20, 36, 74)         # #14244A Rich Navy Card
    ACCENT_CYAN = RGBColor(0, 180, 216)    # #00B4D8 Cyan Accent
    ACCENT_GOLD = RGBColor(226, 192, 68)   # #E2C044 Gold Accent
    TEXT_WHITE = RGBColor(255, 255, 255)   # #FFFFFF Pure White
    TEXT_MUTED = RGBColor(203, 213, 225)   # #CBD5E1 Muted Slate
    TEXT_ACCENT = RGBColor(72, 202, 228)   # #48CAE4 Light Cyan
    BORDER_COLOR = RGBColor(33, 57, 108)   # #21396C Card border

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="المسار 04: أدوات المعرفة والتحقق لتمكين المعرّفين بالإسلام"):
        # Top banner category
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = category_text
        p.alignment = PP_ALIGN.RIGHT
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.font.name = "Arial"

        # Main Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.8))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.alignment = PP_ALIGN.RIGHT
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Arial"

        # Decorative line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = ACCENT_CYAN
        line.line.fill.background()

    def add_card(slide, left, top, width, height, title, body_lines, accent_color=ACCENT_CYAN, title_color=TEXT_WHITE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP

        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.RIGHT
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.font.name = "Arial"

        for line in body_lines:
            p2 = tf.add_paragraph()
            p2.text = f"• {line}"
            p2.alignment = PP_ALIGN.RIGHT
            p2.font.size = Pt(14)
            p2.font.color.rgb = TEXT_MUTED
            p2.font.name = "Arial"
            p2.space_before = Pt(8)

    # ==================== SLIDE 1: COVER ====================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    # Decorative top badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.66), Inches(1.2), Inches(6.0), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = CARD_BG
    badge.line.color.rgb = ACCENT_CYAN
    badge.line.width = Pt(1)
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي 2026م"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.size = Pt(14)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_CYAN
    p_b.font.name = "Arial"

    # Main Project Title
    tb_c = s1.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.333), Inches(1.8))
    p_c = tb_c.text_frame.paragraphs[0]
    p_c.text = "شجرة الحديث الذكية"
    p_c.alignment = PP_ALIGN.CENTER
    p_c.font.size = Pt(44)
    p_c.font.bold = True
    p_c.font.color.rgb = TEXT_WHITE
    p_c.font.name = "Arial"

    p_c2 = tb_c.text_frame.add_paragraph()
    p_c2.text = "Smart Hadith Tree"
    p_c2.alignment = PP_ALIGN.CENTER
    p_c2.font.size = Pt(24)
    p_c2.font.bold = True
    p_c2.font.color.rgb = ACCENT_GOLD
    p_c2.font.name = "Arial"

    # Subtitle / Tagline
    tb_sub = s1.shapes.add_textbox(Inches(1.5), Inches(4.1), Inches(10.333), Inches(1.0))
    p_sub = tb_sub.text_frame.paragraphs[0]
    p_sub.text = "منصة بيانية واستدلالية مدعومة بالذكاء الاصطناعي لكشف انقطاع الأسانيد،\nتمثيل شبكات الرواة بصرياً، وصياغة أحكام الجرح والتعديل الموثقة"
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.font.name = "Arial"

    # Footer Cards
    foot1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(5.6), Inches(4.5), Inches(1.1))
    foot1.fill.solid()
    foot1.fill.fore_color.rgb = CARD_BG
    foot1.line.color.rgb = BORDER_COLOR
    p_f1 = foot1.text_frame.paragraphs[0]
    p_f1.text = "المسار المختار"
    p_f1.alignment = PP_ALIGN.CENTER
    p_f1.font.size = Pt(13)
    p_f1.font.color.rgb = ACCENT_CYAN
    p_f1_2 = foot1.text_frame.add_paragraph()
    p_f1_2.text = "أدوات المعرفة والتحقق لتمكين المعرّفين بالإسلام"
    p_f1_2.alignment = PP_ALIGN.CENTER
    p_f1_2.font.size = Pt(14)
    p_f1_2.font.bold = True
    p_f1_2.font.color.rgb = TEXT_WHITE

    foot2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(5.6), Inches(4.5), Inches(1.1))
    foot2.fill.solid()
    foot2.fill.fore_color.rgb = CARD_BG
    foot2.line.color.rgb = BORDER_COLOR
    p_f2 = foot2.text_frame.paragraphs[0]
    p_f2.text = "الجهة المنظمة"
    p_f2.alignment = PP_ALIGN.CENTER
    p_f2.font.size = Pt(13)
    p_f2.font.color.rgb = ACCENT_GOLD
    p_f2_2 = foot2.text_frame.add_paragraph()
    p_f2_2.text = "مؤسسة باذل الأهلية بالشراكة مع سدايا (SDAIA)"
    p_f2_2.alignment = PP_ALIGN.CENTER
    p_f2_2.font.size = Pt(14)
    p_f2_2.font.bold = True
    p_f2_2.font.color.rgb = TEXT_WHITE

    # ==================== SLIDE 2: THE PROBLEM ====================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "المشكلة والاحتياج الواقعي في دراسة الأسانيد")

    add_card(s2, Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3),
             "تشتت الروايات وتضخم التراجم",
             ["الأسانيد موزعة بين مئات المصنفات وأمهات كتب الحديث والعلل.",
              "تراجم الرواة تتجاوز 100 ألف راوٍ مبعثرين في مجلدات مطولة.",
              "صعوبة الربط التلقائي بين الشيخ وتلميذه عبر مسارات الرواية."],
             title_color=ACCENT_GOLD)

    add_card(s2, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3),
             "بطء التحقق من الاتصال والانقطاع",
             ["يستغرق الباحث ساعات طويلة لمقارنة سنة وفاة الشيخ بميلاد التلميذ.",
              "احتمالية وقوع السهو البشري في كشف الانقطاع الشكلي والزمني.",
              "غياب أدوات رقمية تفحص إمكانية اللقاء والمعاصرة آلياً."],
             title_color=ACCENT_GOLD)

    add_card(s2, Inches(6.9), Inches(4.6), Inches(5.6), Inches(2.3),
             "تعارض أقوال الجرح والتعديل",
             ["اختلاف أحكام الأئمة (كابن حجر، الذهبي، والدارقطني) على الراوي الواحد.",
              "صعوبة تلخيص وتركيب حكم نقدي متوازن لغير المتخصص المتمرس.",
              "خطر هلوسة أدوات الذكاء الاصطناعي العامة عند سؤالها عن رجال الحديث."],
             title_color=ACCENT_GOLD)

    add_card(s2, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.3),
             "غياب التمثيل البصري والبياني",
             ["سلاسل السند تُقرأ كنصوص سردية مصمتة يصعب تخيل تشابكها.",
              "انعدام رؤية 'مدارات الحديث' وتفرعات الطرق والمتابعات بصرياً.",
              "حاجة الدعاة والمعرّفين بالإسلام لأداة بصرية سريعة الإقناع والإيضاح."],
             title_color=ACCENT_GOLD)

    # ==================== SLIDE 3: THE SOLUTION ====================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "الحل المبتكر: شجرة بيانية استدلالية ذكية")

    add_card(s3, Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3),
             "تمثيل بياني تفاعلي (Knowledge Graph)",
             ["تحويل سند الحديث من نص سردي إلى رسم شجري تفاعلي فائق الدقة.",
              "توزيع هرمي منظم باستخدام خوارزميات التخطيط المتطورة (ELK Engine).",
              "تمييز فوري لطبقات الرواة وحالتهم التوثيقية بنظام ألوان دلالي."],
             title_color=ACCENT_CYAN)

    add_card(s3, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3),
             "كشف الانقطاع الزمني الآلي (Inqita' Detection)",
             ["محرك فحص زمني يقارن سنة ميلاد التلميذ بسنة وفاة الشيخ آلياً.",
              "إبراز الانقطاع بخط أحمر متقطع وشارة تحذيرية برهانها مثبت زمنياً.",
              "كشف الانقطاع الشكلي والاستحالة التاريخية في أجزاء من الثانية."],
             title_color=ACCENT_CYAN)

    add_card(s3, Inches(6.9), Inches(4.6), Inches(5.6), Inches(2.3),
             "تلخيص الجرح والتعديل بالذكاء الاصطناعي (RAG)",
             ["توليد معزز بالاسترجاع (RAG) يستند فقط لنصوص الأئمة التاريخية الموثقة.",
              "صياغة خلاصة نقدية متوازنة لمرتبة الراوي بنقرة زر ودون أي هلوسة.",
              "سياسة امتناع صارمة تحيل للمختص عند غياب الدليل أو تعذر الترجيح."],
             title_color=ACCENT_CYAN)

    add_card(s3, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.3),
             "عزل مسارات الضعف والتصدير الفائق",
             ["ميزة الفلترة الذكية لتعتيم الرواة الثقات وإبراز مواطن الضعف في السند.",
              "تصدير الشجرة التفاعلية كصورة عالية الدقة لدعم البحوث والعروض.",
              "بحث فوري بالمتن والراوي عبر 18 مصنفاً حديثياً معتمداً."],
             title_color=ACCENT_CYAN)

    # ==================== SLIDE 4: ARCHITECTURE ====================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "المعمارية التقنية المتكاملة (Clean Architecture)")

    add_card(s4, Inches(9.0), Inches(2.0), Inches(3.6), Inches(4.9),
             "1. طبقة البيانات الضخمة",
             ["موسوعة إتقان (Itqan Dataset) المعتمدة علمياً.",
              "أكثر من 112,800 حديث نبوي مسند.",
              "أكثر من 115,700 راوٍ موثق التواريخ والطبقة.",
              "أكثر من 225,800 رابط إسنادي بين الشيوخ والتلاميذ.",
              "106,000 توثيق جرح وتعديل من كتب التراجم المعتمدة."],
             title_color=ACCENT_GOLD)

    add_card(s4, Inches(4.9), Inches(2.0), Inches(3.6), Inches(4.9),
             "2. الخادم ومحرك المعالجة",
             [".NET 9 Web API مبني وفق Clean Architecture الصارمة.",
              "استعلامات Recursive CTE فائقة السرعة لبناء مسارات الرواية.",
              "خدمة البحث المتقدم وتطبيع النصوص العربية المعقدة.",
              "خدمة فحص الانقطاع وحساب الفوارق الزمنية بين الطبقات.",
              "EF Core مع فهرسة متقدمة في SQL Server لضمان الأداء."],
             title_color=ACCENT_CYAN)

    add_card(s4, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.9),
             "3. واجهة المستخدم والذكاء الاصطناعي",
             ["Next.js (App Router) مع TypeScript و Tailwind CSS.",
              "محرك React Flow مع خوارزميات التخطيط الهرمي ELK.js.",
              "توليد مدعوم بـ Microsoft Semantic Kernel و Gemini LLM.",
              "مكتبة html-to-image لتصدير الشجرة بدقة Vector/High-Res.",
              "واجهة متجاوبة بالكامل تدعم العربية RTL وسهلة الاستخدام."],
             title_color=ACCENT_CYAN)

    # ==================== SLIDE 5: AI & SCIENTIFIC SAFETY ====================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "توظيف الذكاء الاصطناعي ومحددات الأمان والسلامة العلمية")

    add_card(s5, Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3),
             "التوليد المعزز بالاسترجاع (Domain-Specific RAG)",
             ["لا يعتمد النموذج على معلوماته العامة؛ بل يُحقن ببيانات الراوي المسترجعة حصراً.",
              "يسترجع النظام أقوال أئمة الحديث الكبار (ابن حجر، الذهبي، المزي) حرفياً.",
              "يقوم النموذج بتركيب الأقوال ووزنها بأسلوب نقدي محكم كباحث حديثي متمرس."],
             title_color=ACCENT_CYAN)

    add_card(s5, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3),
             "سياسة الحماية والامتناع الصارم (Zero Hallucination)",
             ["منع توليد أي اسم أو تاريخ أو حكم غير موجود في المستندات المسترجعة.",
              "عند فقدان تاريخ الميلاد أو الوفاة، يمتنع النظام عن الحكم بالانقطاع جزافاً.",
              "يظهر النظام بوضوح: 'تاريخ غير متوفر - يتطلب مراجعة مكتبية متخصصة'."],
             title_color=ACCENT_GOLD)

    add_card(s5, Inches(6.9), Inches(4.6), Inches(5.6), Inches(2.3),
             "مستويات المحتوى الأربعة وضبط التوثيق",
             ["المستوى 1: المتن النوبي الصريح والباب والكتاب والمصدر المصنف.",
              "المستوى 2: شجرة الإسناد ورجال السند وأداة التحمل (حدثنا، أخبرنا، عن).",
              "المستوى 3: أقوال أئمة الجرح والتعديل المنسوبة بدقة لكل إمام.",
              "المستوى 4: الملخص التركيبي والتنبيهات المنهجية المدعومة بالذكاء الاصطناعي."],
             title_color=ACCENT_CYAN)

    add_card(s5, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.3),
             "إسناد المراجع والإحالة للمختصين",
             ["توثيق كامل لكل راوٍ مع رقم معرفه في موسوعة إتقان والكتب المعتمدة.",
              "زر مخصص لإحالة المسألة أو الراوي الشائك إلى الباحث البشري.",
              "الشفافية الكاملة: يمكن للمستخدم مراجعة النصوص الأصلية لأقوال الأئمة بنفسه."],
             title_color=ACCENT_GOLD)

    # ==================== SLIDE 6: GRAPH & ANOMALY ENGINE ====================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "محرك كشف الانقطاع والتحليل البياني للأسانيد")

    add_card(s6, Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3),
             "معادلة الفحص الزمني للرواة",
             ["إذا كانت: (سنة ميلاد التلميذ > سنة وفاة الشيخ) -> يُثبت الانقطاع يقيناً.",
              "إذا كان الفارق الزمني مستحيلاً عادة للتحمل والسماع -> يُرفع مؤشر الشبهة.",
              "معالجة كافة حالات فقدان التواريخ دون افتراض مضلل."],
             title_color=ACCENT_CYAN)

    add_card(s6, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3),
             "التجسيد البصري الذكي للعلة",
             ["رسم خط متقطع عريض باللون الأحمر يربط بين الراويين المنقطعين.",
              "إضافة وسم صريح 'انقطاع' فوق الرابط الإسنادي.",
              "أيقونة تحذير تفاعلية توضح بالتمرير: سبب الانقطاع وفارق السنين بدقة."],
             title_color=ACCENT_CYAN)

    add_card(s6, Inches(6.9), Inches(4.6), Inches(5.6), Inches(2.3),
             "عزل الروابط الضعيفة (Weak Links Isolation)",
             ["مفتاح تحكم يتيح خفض شفافية الرواة الثقات إلى 30% بضغطة زر واحدة.",
              "إبراز الرواة الضعفاء أو المتروكين أو المنقطعين بلون بارز وسط الشجرة.",
              "تمكين المعرّف والباحث من اكتشاف علة الحديث في ثوانٍ معدودة."],
             title_color=ACCENT_GOLD)

    add_card(s6, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.3),
             "أدوات التصدير والمشاركة الأكاديمية",
             ["تصدير كامل مساحة الشجرة التفاعلية بصيغة PNG فائقة الجودة.",
              "ملائمة الرسوم للمشاركة في الأبحاث، المحاضرات، والمحتوى الرقمي للدعوة.",
              "حفظ مقياس الرسم ودقة النصوص العربية دون تشويه."],
             title_color=ACCENT_GOLD)

    # ==================== SLIDE 7: CORPUS & COVERAGE ====================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "قاعدة البيانات والشمولية: 18 مصنفاً مسنداً")

    # Big stat boxes
    stats = [
        ("112,813", "حديث مسند ومفهرس", Inches(9.8), ACCENT_CYAN),
        ("115,735", "راوٍ موثق التراجم", Inches(6.8), ACCENT_GOLD),
        ("225,807", "رابط إسنادي متصل", Inches(3.8), ACCENT_CYAN),
        ("18", "مصنفاً حديثياً كاملاً", Inches(0.8), ACCENT_GOLD),
    ]
    for num, lbl, left, col in stats:
        box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(2.7), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = col
        p_n = box.text_frame.paragraphs[0]
        p_n.text = num
        p_n.alignment = PP_ALIGN.CENTER
        p_n.font.size = Pt(28)
        p_n.font.bold = True
        p_n.font.color.rgb = col
        p_l = box.text_frame.add_paragraph()
        p_l.text = lbl
        p_l.alignment = PP_ALIGN.CENTER
        p_l.font.size = Pt(13)
        p_l.font.color.rgb = TEXT_WHITE

    add_card(s7, Inches(6.9), Inches(3.6), Inches(5.6), Inches(3.3),
             "كتب السنة وأمهات المصنفات المدمجة",
             ["صحيح البخاري وصحيح مسلم.",
              "السنن الأربعة: أبو داود، الترمذي، النسائي، ابن ماجه.",
              "مسند الإمام أحمد بن حنبل (26,539 حديثاً).",
              "مصنف ابن أبي شيبة (37,943 حديثاً).",
              "موطأ مالك، سنن الدارمي، الأدب المفرد، وبلوغ المرام."],
             title_color=ACCENT_CYAN)

    add_card(s7, Inches(0.8), Inches(3.6), Inches(5.6), Inches(3.3),
             "توحيد الرواة وربط الأسانيد",
             ["توحيد هوية الرواة (Disambiguation) لمنع تكرار الراوي بأسماء مختلفة.",
              "ربط كل حديث بتبويبه الأصلي وكتابه ورقمه التسلسلي المعتمد عالمياً.",
              "استيعاب 106,000 توثيق جرح وتعديل لمختلف أئمة النقد عبر العصور.",
              "جاهزية كاملة للاستعلام السريع دون الاعتماد على خدمات خارجية معطلة."],
             title_color=ACCENT_GOLD)

    # ==================== SLIDE 8: HACKATHON PLAN ====================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "خطة التطوير خلال أيام التحدي (4 - 6 أكتوبر 2026)")

    add_card(s8, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.9),
             "اليوم الأول (4 أكتوبر): التدقيق الزمني",
             ["تحسين خوارزمية كشف الانقطاع الخفي بالربط مع صيغ التحمل والأداء (عنعنة المدلسين).",
              "تسريع استعلامات الـ Recursive CTE للشجرات الإسنادية الضخمة متفرعة الطرق.",
              "إجراء جلسات الإرشاد الأولى وضبط المتطلبات التوثيقية مع مرشدي التحدي.",
              "إجراء اختبارات دقة على 50 إسناداً معقداً ومقارنتها بأحكام أئمة العلل."],
             title_color=ACCENT_CYAN)

    add_card(s8, Inches(4.8), Inches(2.0), Inches(3.7), Inches(4.9),
             "اليوم الثاني (5 أكتوبر): دمج الطرق والشواهد",
             ["بناء ميزة 'دمج أسانيد الحديث الواحد' لإظهار المتابعات والشواهد في شجرة موحدة.",
              "تطوير لوحة مقارنة الرواة لإبراز التباين والاتفاق في صيغ الرواية.",
              "تحسين صياغة الـ RAG لتسليط الضوء على الرواة المتكلم فيهم بوضوح أكبر.",
              "اختبارات واجهة المستخدم والتأكد من سهولة التنقل والاستجابة على كافة الشاشات."],
             title_color=ACCENT_GOLD)

    add_card(s8, Inches(0.8), Inches(2.0), Inches(3.7), Inches(4.9),
             "اليوم الثالث (6 أكتوبر): التسليم والنشر",
             ["نشر وتأمين النسخة التجريبية الحية (Live Demo) على بيئة سحابية مستقرة.",
              "إنتاج الفيديو التعريفي التوضيحي (أقل من دقيقتين) شاملاً تجربة حية للنظام.",
              "تنظيف وتوثيق مستودع GitHub العام وإرفاق ملفات README وإرشادات التشغيل.",
              "المراجعة الشاملة لمتطلبات التسليم واستكمال الرفع قبل موعد 11:59 مساءً."],
             title_color=ACCENT_CYAN)

    # ==================== SLIDE 9: IMPACT & SUSTAINABILITY ====================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "الأثر المتوقع واستدامة التشغيل والجاهزية")

    add_card(s9, Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3),
             "الأثر على الباحثين والمعرّفين بالإسلام",
             ["توفير أكثر من 90% من الوقت المستغرق في تحقيق صحة السند واتصاله.",
              "تمكين المحاور الإسلامي من إظهار قوة ونظافة الإسناد بالدليل المرئي الحاسم.",
              "نقل علم الإسناد الشريف من حيز التجريد النظري إلى الإدراك البصري المعاصر."],
             title_color=ACCENT_GOLD)

    add_card(s9, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3),
             "قابلية التطبيق والتوسع (Scalability)",
             ["هندسة برمجية حديثة تدعم الحاويات (Docker Containers) ونشر Kubernetes السهل.",
              "قاعدة بيانات محلية متكاملة تضمن سرعة الاستجابة بأقل من 300 ملي ثانية.",
              "بنية مفتوحة تسمح بإضافة كتب الأجزاء والمستخرجات ومعاجم الصحابة بسلاسة."],
             title_color=ACCENT_CYAN)

    add_card(s9, Inches(6.9), Inches(4.6), Inches(5.6), Inches(2.3),
             "تكاليف التشغيل المنخفضة والاستدامة",
             ["الاعتماد على قواعد بيانات مهيكلة يقلل استهلاك استدعاءات نماذج الذكاء الاصطناعي.",
              "حصر استهلاك الذكاء الاصطناعي على التلخيص الدقيق والمنضبط مما يخفض التكاليف.",
              "إمكانية العمل على خوادم سحابية منخفضة التكلفة (Cloud Hosting)."],
             title_color=ACCENT_CYAN)

    add_card(s9, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.3),
             "إتاحة واجهات برمجية مفتوحة (Open API)",
             ["توفير RESTful APIs تتيح للمنصات والمواقع الإسلامية جلب شجرة السند تلقائياً.",
              "تمكين تطبيقات الهواتف الذكية الدعوية من تضمين الرسم الشجري بلمسة زر.",
              "تحويل المنصة إلى بنية تحتية معرفية تخدم مئات المشاريع التقنية الإسلامية."],
             title_color=ACCENT_GOLD)

    # ==================== SLIDE 10: CONCLUSION & TEAM ====================
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10)
    add_header(s10, "فريق العمل والخاتمة")

    add_card(s10, Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3),
             "تكامل كفاءات الفريق وتغطية المهام",
             ["هندسة الذكاء الاصطناعي: ضبط نماذج اللغة وهندسة الـ RAG والاسترجاع الدقيق.",
              "تطوير النظم والواجهات: بناء خوادم .NET 9 ومعمارية React Flow و ELK.",
              "التحقيق الشرعي والمعرفي: ضبط مناهج المحدثين وقواعد الجرح وتدقيق البيانات.",
              "تصميم تجربة المستخدم: تبسيط التعامل مع الرسوم المعقدة وإتاحة الشاشات."],
             title_color=ACCENT_CYAN)

    add_card(s10, Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3),
             "الالتزام بضوابط ومعايير التحدي",
             ["كود مصدري منظم ومفتوح بالكامل على GitHub مع رخصة واضحة وتوثيق تشغيل.",
              "عرض تجريبي مباشر (Live Demo) يعمل فعلياً ومتاح لاختبار لجان التحكيم.",
              "التزام تام بالأمانة العلمية وسلامة الإسناد دون اختلاق أو مبالغة."],
             title_color=ACCENT_CYAN)

    # Final Quote Box
    qbox = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.7), Inches(10.333), Inches(2.1))
    qbox.fill.solid()
    qbox.fill.fore_color.rgb = CARD_BG
    qbox.line.color.rgb = ACCENT_GOLD
    qbox.line.width = Pt(2)

    p_q1 = qbox.text_frame.paragraphs[0]
    p_q1.text = "«الإسناد من الدين، ولولا الإسناد لقال من شاء ما شاء» — عبد الله بن المبارك"
    p_q1.alignment = PP_ALIGN.CENTER
    p_q1.font.size = Pt(18)
    p_q1.font.bold = True
    p_q1.font.color.rgb = ACCENT_GOLD
    p_q1.font.name = "Arial"

    p_q2 = qbox.text_frame.add_paragraph()
    p_q2.text = "مشروع شجرة الحديث الذكية: إعادة إحياء علم الإسناد الخالد بأحدث تقنيات الرسوم البيانية والذكاء الاصطناعي المسند لخدمة أقدس محتوى بشري."
    p_q2.alignment = PP_ALIGN.CENTER
    p_q2.font.size = Pt(15)
    p_q2.font.color.rgb = TEXT_WHITE
    p_q2.font.name = "Arial"
    p_q2.space_before = Pt(12)

    prs.save(output_pptx)
    print(f"Presentation saved successfully to {output_pptx}")

if __name__ == "__main__":
    out_path = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else "Smart_Hadith_Tree_Presentation.pptx")
    create_deck(out_path)
